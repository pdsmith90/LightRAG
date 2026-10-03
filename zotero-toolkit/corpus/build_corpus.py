#!/usr/bin/env python3
"""
build_corpus.py — Turn a Zotero storage/ tree into a clean, metadata-enriched
Markdown corpus ready for LightRAG ingestion.

Strategy (hybrid):
  * PDFs with a real text layer  -> PyMuPDF4LLM (proper reading order; far better
    than `pdftotext -layout` output, which interleaves multi-column text).
  * Image-only / scanned PDFs     -> OCR fallback, in this order:
        1. reuse <--ocr-cache-dir>/<stem>.pdf.txt, if that option is given
        2. run `ocrmypdf` if installed, then PyMuPDF4LLM
        3. reuse <--fallback-md-dir>/<stem>.md as a last resort, if given
  * epub / docx / pptx / djvu / html / txt -> best available extractor
    (HTML snapshots prefer Zotero's own .zotero-ft-cache text).
  * Every document gets a metadata header (title / authors / year / DOI /
    publication / tags / Zotero link) prepended from zotero_metadata.json.

Output file naming:  <STORAGE_KEY>__<slug>.md
  The storage key prefix guarantees uniqueness AND lets answers cite back to
  Zotero via  zotero://select/library/items/<STORAGE_KEY>.

Resumable & parallel. Safe to re-run; only (re)builds missing/outdated outputs.

Usage:
  python3 build_corpus.py \
      --zotero ~/Zotero \
      --out    ~/Zotero/rag_corpus \
      --meta   ~/Zotero/zotero_metadata.json \
      --jobs   8 --engine pymupdf4llm --ocr auto

Run  `python3 build_corpus.py --help`  for all options.
"""

from __future__ import annotations
import argparse
import glob
import json
import os
import re
import time
import shutil
import subprocess
import tempfile
from concurrent.futures import ProcessPoolExecutor, as_completed

# Denylist semantics are shared with quarantine_junk.py, which runs immediately
# after this script and would move any denylisted output straight back out.
# Import rather than reimplement so the two can never drift. Degrade to an empty
# denylist if that module is missing: rebuilding a denylisted file is wasteful,
# but failing the whole corpus build over it would be far worse.
try:
    from quarantine_junk import load_denylist
except Exception as _deny_err:
    print(
        f"[build_corpus] WARN: denylist unavailable ({_deny_err}); skipping nothing",
        flush=True,
    )

    def load_denylist(path):
        return set()


# ---------- source-file selection ----------
PDF, EPUB, DOCX, PPTX, DJVU, HTML, TXT = "pdf epub docx pptx djvu html txt".split()
PRIORITY = [
    PDF,
    EPUB,
    DOCX,
    PPTX,
    DJVU,
    HTML,
    TXT,
]  # which attachment to prefer per folder
EXT2KIND = {
    ".pdf": PDF,
    ".epub": EPUB,
    ".docx": DOCX,
    ".pptx": PPTX,
    ".ppt": PPTX,
    ".djvu": DJVU,
    ".html": HTML,
    ".htm": HTML,
    ".txt": TXT,
}

_slug_re = re.compile(r"[^a-zA-Z0-9._-]+")


def slug(s: str, n: int = 80) -> str:
    s = _slug_re.sub("_", (s or "").strip()).strip("_")
    return s[:n] or "untitled"


# ---------- text cleanup ----------
_pic_re = re.compile(
    r"<!--\s*Start of picture text\s*-->.*?<!--\s*End of picture text\s*-->",
    re.DOTALL | re.IGNORECASE,
)
_manyblank = re.compile(r"\n{3,}")
_datauri_re = re.compile(r"!?\[[^\]]*\]\(data:[^)]*\)")
# 200, not 800: an 800-char bar let 200-799-char base64 runs through into
# indexed documents, and a 200-char unbroken [A-Za-z0-9+/] run never occurs in
# real prose (quarantine_junk.py uses the same 200 for its ratio).
_b64_re = re.compile(r"[A-Za-z0-9+/]{200,}={0,2}")
# HTML-tag residue: pandoc -t gfm passes raw HTML through, so saved web pages
# leave <span>/<div>/<a ...> markup behind. Strip the tag
# markup only, never element text. The (\s[^<>]*)? form requires a real tag shape
# (`<b>` or `<a href=...>`), so math like "a<b, and c>d" is never eaten.
_script_style_re = re.compile(
    r"<(script|style)\b[^>]*>.*?</\1\s*>", re.DOTALL | re.IGNORECASE
)
_tag_re = re.compile(
    r"</?(?:a|abbr|article|aside|b|body|button|center|div|em|figcaption|figure|font"
    r"|footer|form|head|header|html|i|iframe|img|input|label|li|link|main|meta|nav"
    r"|noscript|ol|p|path|section|small|source|span|strong|sub|sup|svg|table|tbody"
    r"|td|th|thead|tr|u|ul|video)\b(?:\s[^<>]*)?/?>",
    re.IGNORECASE,
)
# ---- page furniture / conversion residue -------------------------------------
# In a corpus of converted papers, repeated lines were a few percent of the text
# but touched most files, and some indexed chunks were almost all boilerplate
# (a retrieval hit on one wastes most of its context budget).
#
# Rule - lines repeating within one document (running headers, per-page licence
# stamps). KEEPING THE FIRST OCCURRENCE is what makes this safe: a textbook
# chapter title appears in the TOC and as a running header on every page, and
# dropping every copy would delete the heading itself. Structural markdown is
# excluded outright so tables, lists, quotes and headings are never candidates.
_struct_re = re.compile(r"^\s*(\||#{1,6}\s|[-*+]\s|>\s|```|\d+\.\s)")
_REPEAT_MIN = 5  # occurrences before a line counts as furniture
_REPEAT_MINLEN = 40  # chars; shorter lines are too ambiguous to judge
_REPEAT_MAX_SHRINK = 0.25  # bail out above this; see SAFETY VALVE below


def _drop_repeats(md: str) -> str:
    lines = md.split("\n")

    def norm(t):
        return re.sub(r"\d+", "#", re.sub(r"\s+", " ", t.strip()))

    counts = {}
    for ln in lines:
        t = ln.strip()
        if len(t) >= _REPEAT_MINLEN and not _struct_re.match(ln):
            k = norm(t)
            counts[k] = counts.get(k, 0) + 1
    kill = {k for k, n in counts.items() if n >= _REPEAT_MIN}
    if not kill:
        return md
    seen, out = {}, []
    for ln in lines:
        t = ln.strip()
        if len(t) >= _REPEAT_MINLEN and not _struct_re.match(ln):
            k = norm(t)
            if k in kill:
                seen[k] = seen.get(k, 0) + 1
                if seen[k] > 1:
                    continue  # keep the first, drop later repeats
        out.append(ln)
    cleaned = "\n".join(out)
    # SAFETY VALVE. A real paper sheds a few percent of running headers. Anything
    # shedding more than this is not a paper with furniture -- it is an HTML
    # snapshot (e.g. a saved PDF.js viewer page, where every text run is wrapped
    # in its own <span> so "markup" lines carry the actual words, and dropping
    # duplicates silently eats characters). Those belong to quarantine_junk.py,
    # which moves whole files; never let a line rule mangle one. Such a
    # snapshot can lose about half its characters to this rule.
    if md and (len(md) - len(cleaned)) / len(md) > _REPEAT_MAX_SHRINK:
        return md
    return cleaned


# ---- reference-list sections ---------------------------------------------------
# A bibliography is keyword-dense in exactly the words a query uses -- paper
# titles -- so its chunks win the reranker while carrying no content. In a
# knowledge base of academic papers with about a hundred thousand chunks, more than
# half of the chunks returned for a typical question were reference-list text, a few
# percent of all chunks held >=5 "doi:" strings, and the extraction LLM minted one
# author-name entity per citation ("A person involved in..." entities, a sizeable
# share of the graph). Cut the lists at the source.
#
# Detection is per LINE, not per heading: pymupdf4llm renders the heading 25+
# ways ("## **References**", bare "REFERENCES", "_Bibliography_", any depth to
# ######) and textbooks carry one list per chapter. A line is a reference ENTRY
# when it has a year AND a citation cue (doi/url, volume(issue)/pages, "pp."/
# "vol.", "In: Editor", "(Eds.)", or a "Surname, I. I." author token). A run of
# >= _REF_MIN_ENTRIES such lines -- tolerating <= _REF_MAX_GAP consecutive
# non-matching lines, for wrapped entries and running headers -- is removed from
# its first entry to its last (plus trailing wrapped lines that still carry a
# doi/url/pages cue), together with a "References"-style heading directly above
# it. Table rows and code fences never join a run and always end one, so a data
# table listing years and DOIs survives. A run shorter than the minimum is kept:
# the leak is small, and prose with five consecutive year+cue lines does not
# occur (a paragraph is a single line here).
#
# An entry must also START like one -- "Surname, I.", "SURNAME W. A.", "Smith BD,",
# "I. Surname", after an optional "-", "[12]" or "12." marker. Without that test
# a sweep over a few thousand parsed documents flagged OCR'd journal scans whose
# pymupdf4llm output is one line per PAGE: a page carrying a year and
# "VOL. 323" counted as an entry, six pages made a run, and only the shrink
# valve saved the paper (99% of it would have gone). In a two-page comment the
# same logic reached back into the metadata header. Organisation-authored
# entries ("WHO (2013), ...") fail the start test and ride along as tolerated
# gap lines instead; that is the accepted recall cost.
#
# v3 of the detector. Books, handbooks, theses and documents whose indexed chunks
# still held reference lists were untouched by v2, for six reasons found by
# sampling them: (1) Springer chapter lists carry a bullet AND an index
# ("- 1.1 J.R. Smith: ..."), v2's marker allowed one or the other; (2) markdown
# emphasis and HTML tags split the cues ("**12** (3), 101-118", "<span ...>Smith,
# S.,"), so every line is now judged on a NORMALISED copy with emphasis, tags and
# pymupdf's detached accents ("Andr´e") removed; (3) Springer-basic authors
# ("Smith R (1984)", "Jones RA, Brown SM.") and Chicago full first names
# ("Smith, Martin. 1998.") matched neither the start test nor the author cue;
# (4) hard-wrapped output (pandoc/EndNote, OCR'd two-column scans) puts the year
# two or three lines below the author, so an entry may now span up to
# _REF_MAX_WRAP continuation lines that do not themselves start like an entry;
# (5) line-numbered manuscripts ("513 Smith, R.M., ...") and same-author dashes
# ("-- , Jones, T.,") needed marker forms; (6) "vol:first-last" and
# "first-last (year)" page cues were missing. The output text is never
# normalised -- only the decision is.
_REF_MIN_ENTRIES = 5
_REF_MAX_GAP = 2
_REF_MAX_WRAP = 4  # continuation lines one wrapped entry may span (v3)
# SAFETY VALVE: above this the document IS a bibliography (or the detector is
# wrong) -- leave it whole for quarantine_junk.py / a human. Short letters
# run ~35% references, review papers ~30%.
_REF_MAX_SHRINK = 0.60
# Second tier: every valve trip is logged, and a real body relaxes it. A review article
# whose list came to just over 0.60 was ingested with its whole list, and most of its chunks
# were references. When a real body survives (>= _REF_BIG_BODY chars) a list of up to 75% is
# still a list; the documents the valve exists for leave ~1-2 KB (misfires removed roughly
# 80-100% of the text, leaving <2 KB -- in a sweep no document between 0.60 and 0.75 had a
# body that large, so this changes edge cases only).
_REF_MAX_SHRINK_BIG = 0.75
_REF_BIG_BODY = 10_000
_ref_doc = (
    ""  # set by convert() so a valve decision can name its document in the build log
)
_ref_logged = set()  # clean_md runs twice on PDFs: log each document once
_ref_heading_re = re.compile(
    r"^\s*(?:#{1,6}\s*)?(?:\*\*|__|_|\*){0,2}\s*(?:\d{1,4}(?:\.\d{1,2})*\.?\s*)?(?:\*\*|__|_|\*){0,2}\s*"
    r"(?:references?(?:\s+(?:and\s+notes|cited|list))?|bibliography|bibliographie|r[ée]f[ée]rences"
    r"|literature(?:\s+cited)?|literaturverzeichnis|works\s+cited|sources?\s+cited|notes\s+and\s+references"
    r"|(?:further|selected|suggested|recommended|additional)\s+(?:readings?|references|literature))"
    r"\s*:?\s*\d{0,4}\s*(?:\*\*|__|_|\*){0,2}\s*#*\s*$",
    re.IGNORECASE,
)
_ref_year_re = re.compile(r"(?<!\d)(?:1[6-9]\d{2}|20[0-3]\d)[a-z]?(?!\d)")
_ref_strong_cue_re = re.compile(
    r"\bdoi\b|doi\.org|https?://|\barxiv\b|\bpp?\.\s*\d|\bvol\.?\s*\d"
    r"|\d{1,4}\s*\(\s*[A-Za-z]?\d{1,4}\s*\)\s*[,:]?\s*[A-Za-z]?\d"  # 12(3), 1001  /  88(D4), 12,345
    r"|\b\d{1,4}\s*,\s*\d{1,5}\s*[–—-]\s*\d{1,5}\b"  # 62, 1001–1010
    r"|\b\d{1,4}\s*:\s*[A-Za-z]?\d{1,5}\s*[–—-]\s*\d{1,5}\b"  # 45:101–112  (v3)
    r"|\b\d{1,5}\s*[–—-]\s*\d{1,5}\s*\(\s*(?:1[6-9]\d{2}|20[0-3]\d)\s*\)",  # 101–118 (2007)  (v3, Springer/Nature)
    re.IGNORECASE,
)
_ref_weak_cue_re = re.compile(
    r"\bIn:\s*[A-Z]|\(eds?\.?\)|\bedited\s+by\b"
    r"|\b(?:ph\.?\s?d\.?|doctoral|master'?s)\s+(?:thesis|dissertation)\b|\bdissertation\b"  # (v3) books/theses
    r"|\b\d+(?:st|nd|rd|th)\s+edn?\b|\bedn\.|\btech(?:nical)?\.?\s+(?:rep(?:ort)?|note|memo)",
    re.IGNORECASE,
)
_U = r"[A-ZÀ-ÖØ-Þ]"  # a capital, Latin-1 included (Özdemir, Étienne)
_L = r"(?:[^\W\d_]|['’-])"  # a letter in any script, or a name hyphen/apostrophe
_ref_author_re = re.compile(  # case matters
    r"\b" + _U + _L + r"+,\s*(?:[A-Z]\.[\s-]*){1,3}"  # Surname, I. I.  (Müller, J.)
    r"|\b"
    + _U
    + _L
    + r"+\s+(?:[A-Z]\.\s*){1,3}(?=[&,(\d]|and\b)"  # SURNAME W. A. & / Smith W. M. 1966 (books: no pages/doi)
    r"|\b"
    + _U
    + _L
    + r"+\s+[A-Z]{1,3},\s+"
    + _U
    + _L
    + r"+\s+[A-Z]{1,3}\b"  # Smith BD, Jones S        (v3, Springer basic)
    r"|\b"
    + _U
    + _L
    + r"+\s+[A-Z]{1,3}\s*\(\s*(?:1[6-9]\d{2}|20[0-3]\d)[a-z]?\s*\)"  # Smith R (1984)           (v3)
    r"|\b"
    + _U
    + _L
    + r"+,\s+(?:[A-Z][a-z]+(?:\s+[A-Z]\.)?|[A-Z]{1,3})(?:\.|,|\s+(?:and|&)\b)"  # Smith, Martin. / Doe, JR. (v3, Chicago)
    r"|(?:\b[A-Z]\.\s?){1,3}" + _U + _L + r"+\s*(?:[:,]|\s(?:and|&)\b)"
)  # J.R. Smith: / J. Jones,  (v3, Springer chapter style)
_ref_struct_re = re.compile(r"^\s*(?:\||```)")
_REF_PARTICLES = r"(?:[Vv]an|[Vv]on|[Dd]e|[Dd]er|[Dd]en|[Dd]el|[Dd]i|[Dd]a|[Ll]e|[Ll]a|[Dd]u|[Dd]os|[Dd]as|[Tt]e|[Tt]er|[Aa]f|[Zz]u)"
_ref_start_re = re.compile(
    r"^\s*(?:[^\w\s\[(#]{1,3}\s*,?\s*)?"  # bullet, blockquote, same-author dash "-- ,", OCR junk ("/'Smith")
    r"(?:(?:\[\d{1,4}\]|\(\d{1,4}\)|\d{1,4}(?:\.\d{1,3})*[.)]?)\s*){0,2}"  # [12]  (12)  12.  1.1  513  and "1874 [1]" (line number + index)
    r"(?:" + _REF_PARTICLES + r"\s+)*"  # name particles: van, De, ...
    r"(?:"
    + _U
    + r"[\w'’\-]+(?:\s+(?:"
    + _U
    + r"[\w'’\-]+|"
    + _REF_PARTICLES
    + r"\b)){0,2}\s*,"  # Surname, / Van Buren, / De Vries-van Berg,
    r"|"
    + _U
    + r"[\w'’\-]+\s+(?:(?:[A-Z]\.\s*){1,3}|[A-Z]{1,3}\s*(?:,|\())"  # Surname I. I. / Smith BD, / Doe EG (
    r"|(?:[A-Z]\.\s*){1,3}" + _U + r"[\w'’\-]+)"
)  # I. Surname / J.R. Smith
_ref_url_re = re.compile(r"\bdoi\b|doi\.org|https?://", re.IGNORECASE)
_ref_tag_re = re.compile(r"<[^<>]{1,200}>")
_ref_emph_re = re.compile(r"[*~]+|(?<!\w)_+|_+(?!\w)")
_ref_accent_re = re.compile(
    "\\s?[\u00a8\u00b4\u0060\u02c6\u02dc\u00b8\u02d8\u02c7\u00af]\\s?"
)  # pymupdf detaches accents: Andr´e, Kova ˇc
_ref_md_heading_re = re.compile(r"^\s*#{1,6}\s")
_ref_term_re = re.compile(
    r"(?:[.)\]]|\bdoi:?\s*\S+|doi\.org/\S+|https?://\S+)\s*$", re.IGNORECASE
)  # an entry ends like this


def _ref_norm(line: str) -> str:
    """The copy of a line the detector judges: no HTML tags, no markdown emphasis, no
    detached accents, single spaces. The document text itself is never changed."""
    return re.sub(
        r"[ \t]{2,}",
        " ",
        _ref_accent_re.sub("", _ref_emph_re.sub("", _ref_tag_re.sub("", line))),
    )


def _ref_weight(line: str) -> int:
    # pymupdf4llm sometimes emits an entire bibliography as ONE paragraph (a few
    # percent of parsed documents in one sweep); per-line counting sees a single entry and
    # keeps it. Such a line starts like a reference and packs >= _REF_MIN_ENTRIES
    # years with >= 3 doi/url cues -- prose never does -- so it counts as a run.
    if (
        len(line) > 600
        and len(_ref_url_re.findall(line)) >= 3
        and len(_ref_year_re.findall(line)) >= _REF_MIN_ENTRIES
    ):
        return _REF_MIN_ENTRIES
    return 1


def _ref_body_ok(text: str) -> bool:
    return bool(_ref_year_re.search(text)) and bool(
        _ref_strong_cue_re.search(text)
        or _ref_weak_cue_re.search(text)
        or _ref_author_re.search(text)
    )


def _is_ref_entry(line: str) -> bool:
    nrm = _ref_norm(line)
    return bool(_ref_start_re.match(nrm)) and _ref_body_ok(nrm)


def _ref_entry_end(lines: list, i: int, n: int) -> int:
    """Index of the last line of the reference entry that starts at line i, or -1.
    An entry may run over up to _REF_MAX_WRAP continuation lines: while its text is
    not yet a complete entry, or is one but does not end like one ('.', ')', a DOI/URL)
    -- hard-wrapped pandoc/EndNote and OCR output puts the year and the journal two or
    three lines below the authors. A continuation line must not start like an entry,
    unless the text so far is an unfinished author list ending in a comma."""
    text = _ref_norm(lines[i])
    if not _ref_start_re.match(text):
        return -1
    end, j, blanks = (i if _ref_body_ok(text) else -1), i + 1, 0
    while j < n and j - i <= _REF_MAX_WRAP:
        ln = lines[j]
        if not ln.strip():
            # pymupdf emits OCR'd hard-wrapped lines as separate paragraphs: an entry that
            # is still incomplete may continue after ONE blank line; a complete one ends here
            if end < 0 and blanks == 0 and j + 1 < n and lines[j + 1].strip():
                blanks, j = 1, j + 1
                continue
            break
        if _ref_struct_re.match(ln) or _ref_md_heading_re.match(ln):
            break
        ok = end >= 0
        if ok and _ref_term_re.search(text):
            break  # complete and terminated: the entry ends here
        nrm = _ref_norm(ln)
        if _ref_start_re.match(nrm) and (ok or not text.rstrip().endswith(",")):
            break  # the next entry starts
        text = text + " " + nrm
        if _ref_body_ok(text):
            end = j
        j += 1
    return end


def _strip_references(md: str) -> str:
    lines = md.split("\n")
    n = len(lines)
    spans = []
    entry_lines = set()
    i = 0
    while i < n:
        if _ref_struct_re.match(lines[i]):
            i += 1
            continue
        e = _ref_entry_end(lines, i, n)
        if e < 0:
            i += 1
            continue
        start, last, entries, gap, j = i, e, _ref_weight(_ref_norm(lines[i])), 0, e + 1
        run_lines = set(range(i, e + 1))
        while j < n:
            ln = lines[j]
            if _ref_struct_re.match(ln):
                break
            if ln.strip():  # blank lines are neutral
                e2 = _ref_entry_end(lines, j, n)
                if e2 >= 0:
                    last, entries, gap = e2, entries + _ref_weight(_ref_norm(ln)), 0
                    run_lines.update(range(j, e2 + 1))
                    j = e2 + 1
                    continue
                gap += 1
                if gap > _REF_MAX_GAP:
                    break
            j += 1
        if entries < _REF_MIN_ENTRIES:
            i = max(j, start + 1)  # every line in [start, j) was already judged
            continue
        end = last
        # wrapped tail of the final entry: no year on the line, but a doi/url/pages cue
        k, absorbed = end + 1, 0
        while k < n and absorbed < _REF_MAX_GAP and lines[k].strip():
            if _ref_struct_re.match(lines[k]) or not _ref_strong_cue_re.search(
                _ref_norm(lines[k])
            ):
                break
            end, absorbed, k = k, absorbed + 1, k + 1
        # leading fragment of the first entry (pandoc/EndNote output wraps the author
        # list, so the first line has authors and no year): absorb up to _REF_MAX_GAP
        # non-blank lines directly above that start like an entry AND carry an author cue
        h, absorbed = start - 1, 0
        while h >= 0 and absorbed < _REF_MAX_GAP and lines[h].strip():
            nrm = _ref_norm(lines[h])
            if _ref_struct_re.match(lines[h]) or not (
                _ref_start_re.match(nrm) and _ref_author_re.search(nrm)
            ):
                break
            start, absorbed, h = h, absorbed + 1, h - 1
        h = start - 1
        while h >= 0 and not lines[h].strip():
            h -= 1
        if h >= 0 and _ref_heading_re.match(lines[h]):
            start = h
        spans.append((start, end))
        entry_lines |= run_lines
        i = end + 1
    if not spans:
        return md
    # Gap lines inside a run are dropped with it (wrapped fragments, running headers,
    # organisation-authored entries) -- except a prose paragraph that a two-column layout
    # interleaved into the list (seen: a Conclusions paragraph sat between two
    # entries): long, and carrying neither a citation cue nor an author token. A heading
    # directly above such a paragraph is a real section heading and stays with it.
    drop, keep = set(), set()
    for s, e in spans:
        for idx in range(s, e + 1):
            ln = lines[idx]
            if idx in entry_lines or not ln.strip() or _ref_md_heading_re.match(ln):
                continue
            nrm = _ref_norm(ln)
            if (
                len(nrm) > 200
                and not _ref_strong_cue_re.search(nrm)
                and not _ref_author_re.search(nrm)
            ):
                keep.add(idx)
                h = idx - 1
                while h > s and not lines[h].strip():
                    h -= 1
                if h >= s and _ref_md_heading_re.match(lines[h]):
                    keep.add(h)
        drop.update(idx for idx in range(s, e + 1) if idx not in keep)
    cleaned = "\n".join(ln for idx, ln in enumerate(lines) if idx not in drop)
    shrink = (len(md) - len(cleaned)) / len(md) if md else 0.0
    if shrink > _REF_MAX_SHRINK:
        relaxed = shrink <= _REF_MAX_SHRINK_BIG and len(cleaned) >= _REF_BIG_BODY
        if _ref_doc and _ref_doc not in _ref_logged:
            _ref_logged.add(_ref_doc)
            print(
                f"[build_corpus] VALVE {'relaxed' if relaxed else 'kept whole'} {_ref_doc}: references "
                f"{shrink:.1%} of the text, {len(cleaned)} chars of body"
                f"{' -- stripped' if relaxed else ' -- ingested WITH its reference list'}",
                flush=True,
            )
        if not relaxed:
            return md
    return cleaned


# ---- back-of-book indexes ------------------------------------------------------
# A subject, name or author index is a page-number list of every term in a book.
# The extraction LLM mints one entity per entry with nothing to say about it, and
# an index chunk is where a small extraction model loops: one such chunk can
# generate until its request times out, on every retry, and fail the whole
# document. Cut the index at the source.
#
# A region starts at an index HEADING -- a markdown heading or a bold-only line
# whose text, without markup and page numbers, is an index title ("Subject Index",
# "Index", "Index of Programs", "AUTHOR INDEX", "Sachverzeichnis") -- in the last
# quarter of the document. It runs to the end of the document, or to the first
# heading that is not index furniture (letter groups "A", "Symbols", page running
# headers such as "512 Author Index", further index titles, or a short heading
# whose section still reads as index entries -- an entry with sub-entries rendered
# as a heading), or to a table row: publisher advertisements and back matter such
# as a table of constants follow some indexes and stay. A region is cut only when
# >= _IDX_MIN_DENSITY of its non-blank lines read as index entries (page numbers,
# or a "see" cross-reference) and it holds at least _IDX_MIN_LINES of them. The
# position gate bounds any cut to the last quarter of the document (over a few
# thousand parsed documents the cuts hit books only, the largest about a tenth of
# one), and the output text is never rewritten -- only whole lines are dropped.
_IDX_MIN_POS = 0.75
_IDX_MIN_DENSITY = 0.60
_IDX_MIN_LINES = 5
_idx_markup_re = re.compile(r"<[^<>]{1,40}>|[*_#`]+")
_idx_bold_line_re = re.compile(r"^\s*(?:\*\*|__)(?=\S).{0,80}?(?:\*\*|__)\s*$")
_idx_md_heading_re = re.compile(r"^\s*#{1,6}\s")
_idx_title_re = re.compile(
    r"(?:(?:subject|name|author|keyword|general|notation|symbol|topic|program|cumulative|combined)s?\s+)?"
    r"(?:index|indices|indexes)"
    r"|(?:index|indices|indexes)\s+of\s+[a-z][a-z\s-]{1,40}"
    r"|(?:sach|namens?|stichwort|personen|autoren|schlagwort)(?:-?\s*und\s*(?:sach|namens?))?-?\s*verzeichnis",
    re.IGNORECASE,
)
_idx_furniture_re = re.compile(
    r"[^\W\d_]{1,3}(?:\s*[-–—]\s*[^\W\d_]{1,3})?|\d*|symbols?|numerics?|numbers|numerals|greek(?:\s+letters)?",
    re.IGNORECASE,
)
_idx_page_re = re.compile(
    r"(?<![\w.])\d{1,4}(?:\s*[-–—]\s*\d{1,4})?(?:ff?\.?|[a-z])?(?![\w])"
)
_idx_word_re = re.compile(r"[^\W\d_]{2,}")
_idx_see_re = re.compile(r"\bsee(?:\s+also)?\b", re.IGNORECASE)
_idx_table_re = re.compile(r"^\s*\|")
_idx_logged = set()


def _idx_heading_text(line: str):
    """The text of a heading or bold-only line without markup, else None."""
    if not (_idx_md_heading_re.match(line) or _idx_bold_line_re.match(line)):
        return None
    return re.sub(r"\s+", " ", _idx_markup_re.sub(" ", line)).strip()


def _idx_title(text: str) -> bool:
    """An index title, allowing a page number on either side (running headers)."""
    t = re.sub(r"^\d{1,4}\s+|\s+\d{1,4}$", "", text).strip(" .:")
    return bool(_idx_title_re.fullmatch(t))


def _idx_entry_like(line: str) -> bool:
    t = re.sub(r"\s+", " ", _idx_markup_re.sub(" ", line)).strip()
    if not t:
        return False
    pages = len(_idx_page_re.findall(t))
    if pages == 0:
        return len(t) <= 200 and bool(_idx_see_re.search(t))
    return len(t) <= 160 or pages >= 0.15 * max(len(_idx_word_re.findall(t)), 1)


def _idx_continues(lines: list, j: int) -> bool:
    """A short heading inside an index whose section still reads as index entries --
    an entry with sub-entries rendered as a heading. Back matter after an index does
    not pass: a long title (a publisher's series page), a table under the heading (a
    table of constants), or prose (a list of participants)."""
    if len(_idx_heading_text(lines[j]).split()) > 4:
        return False
    seg = []
    for ln in lines[j + 1 : j + 61]:
        if _idx_table_re.match(ln):
            return False
        if _idx_heading_text(ln) is not None:
            break
        if ln.strip():
            seg.append(ln)
    return len(seg) >= 2 and sum(map(_idx_entry_like, seg)) >= 0.7 * len(seg)


def _idx_region_end(lines: list, start: int) -> int:
    """Index of the first line AFTER the index region that starts at `start`."""
    for j in range(start + 1, len(lines)):
        ln = lines[j]
        if _idx_table_re.match(ln):
            return j
        text = _idx_heading_text(ln)
        if text is None:
            continue
        if (
            _idx_title(text)
            or _idx_furniture_re.fullmatch(re.sub(r"[^\w\s]", "", text).strip())
            or _idx_entry_like(ln)
            or _idx_continues(lines, j)
        ):
            continue  # letter group, running header, bold page number, next index, sub-entry head
        return j
    return len(lines)


def _strip_back_index(md: str) -> str:
    doc = _ref_doc
    lines = md.split("\n")
    n, total = len(lines), len(md)
    if not total:
        return md
    pos, offsets = 0, []
    for ln in lines:
        offsets.append(pos)
        pos += len(ln) + 1
    drop, verdicts, i = [], [], 0
    while i < n:
        text = _idx_heading_text(lines[i])
        if text is None or not _idx_title(text) or offsets[i] < _IDX_MIN_POS * total:
            i += 1
            continue
        end = _idx_region_end(lines, i)
        body = [
            ln
            for ln in lines[i + 1 : end]
            if ln.strip() and _idx_heading_text(ln) is None
        ]
        entries = sum(1 for ln in body if _idx_entry_like(ln))
        density = entries / len(body) if body else 0.0
        if entries >= _IDX_MIN_LINES and density >= _IDX_MIN_DENSITY:
            drop.append((i, end))
            verdicts.append(
                f"'{text[:40]}' line {i} ({offsets[i] / total:.1%}), {end - i} lines, density {density:.2f}"
            )
        else:
            verdicts.append(
                f"KEPT '{text[:40]}' line {i}: {entries} entry lines, density {density:.2f}"
            )
        i = end
    if not drop:
        if doc and verdicts and doc not in _idx_logged:
            _idx_logged.add(doc)
            print(
                f"[build_corpus] INDEX kept {doc}: " + "; ".join(verdicts), flush=True
            )
        return md
    if doc and doc not in _idx_logged:
        cut = sum(offsets[e] if e < n else total for _, e in drop) - sum(
            offsets[s] for s, _ in drop
        )
        _idx_logged.add(doc)
        print(
            f"[build_corpus] INDEX stripped {doc}: {cut} chars ({cut / total:.1%}); "
            + "; ".join(verdicts),
            flush=True,
        )
    keep = [True] * n
    for s, e in drop:
        for k in range(s, e):
            keep[k] = False
    return "\n".join(ln for k, ln in enumerate(lines) if keep[k])


def clean_md(md: str) -> str:
    md = _pic_re.sub("", md)  # drop OCR'd figure-scribble noise
    md = _datauri_re.sub("", md)  # embedded data: images
    md = _b64_re.sub(
        "", md
    )  # base64 blobs dumped as text (saved web pages carry multi-megabyte lines of it)
    md = _script_style_re.sub("", md)  # inline <script>/<style> blocks, content and all
    md = re.sub(r"(?:<br>\s*){2,}", "\n", md)  # collapse <br> soup
    md = re.sub(r"<br\s*/?>", "\n", md, flags=re.IGNORECASE)  # lone <br> residue
    md = _tag_re.sub("", md)  # HTML tag residue from pandoc raw-html passthrough
    md = _strip_back_index(md)  # back-of-book subject/name/author indexes
    md = _strip_references(md)  # bibliography / reference-list sections
    md = _drop_repeats(md)  # running headers, page stamps
    md = _manyblank.sub("\n\n", md)
    return md.strip()


def looks_empty(text: str, min_chars: int) -> bool:
    return len(re.sub(r"\s+", "", text)) < min_chars


# A text layer whose fonts map every glyph into the Unicode Private Use Area (U+E000-F8FF
# and planes 15-16) is not empty, so looks_empty never sent it to OCR: PDF extractors and
# Zotero's own full-text index return the glyph codes as "text", and one scanned
# encyclopedia became thousands of chunks of them. On a corpus of several thousand
# academic PDFs the garbled files scored above 0.75 and every genuine file below 0.05
# (symbol-font Greek in a mathematics handbook was the highest), so 0.30 sits in a clean
# gap. quarantine_junk.py applies the same bar to the corpus root as a second line.
GARBLED_PUA_RATIO = 0.30
_pua_re = re.compile(r"[\ue000-\uf8ff\U000f0000-\U0010ffff]")


def pua_share(text: str) -> float:
    """Private Use Area code points over non-whitespace characters (0.0 under 500 of them)."""
    nonspace = len(re.sub(r"\s+", "", text))
    return len(_pua_re.findall(text)) / nonspace if nonspace >= 500 else 0.0


def looks_garbled(text: str) -> bool:
    return pua_share(text) >= GARBLED_PUA_RATIO


# ---------- pages pymupdf4llm drops: repair them from the PDF's own text layer ----------
# Two failure classes, measured on a corpus of several thousand academic PDFs:
#
# 1. Scans whose invisible OCR text is drawn BEFORE the full-page image (common in
#    publishers' legacy journal scans; page.get_texttrace() shows the text spans at
#    sequence 0-2 with type 3 and the image after them). pymupdf4llm 1.28, built on
#    pymupdf-layout, treats text under a later image as hidden, so such a page comes out
#    as its visible download watermark only: a few hundred characters of a
#    3,000-6,000-character layer, once per page. Scans drawn image-first are fine.
# 2. Born-digital papers lose whole pages that carry a figure: with its default
#    use_ocr=True the engine OCRs such a page and keeps the Tesseract text instead of
#    the real layer (one 14-page paper kept 6-39 % on 7 pages). Not caught by any
#    document-level size check, because the document as a whole keeps 50-70 %.
#
# page.get_text() reads the layer in both cases. Every PDF therefore goes through one
# per-page pass (page_chunks=True; joined with "" it is byte-identical to the plain
# output) with the engine's own OCR off -- it helps neither class, costs about a second
# per page, and on ordinary pages it swapped the real text for a worse OCR of the page
# image -- and a page that kept less than TEXTLAYER_PAGE_KEEP of a text layer of at
# least TEXTLAYER_PAGE_MIN characters is replaced by that layer as paragraphs. Pages
# the engine handled keep their markdown. Image-only PDFs (full-page images, no text
# layer at all) return "" so convert() takes its OCR route; before, a visible watermark
# alone kept such a scan from looking empty.
TEXTLAYER_PAGE_KEEP = 0.4  # a page that kept less than this of its layer is repaired
TEXTLAYER_PAGE_MIN = 400  # ...if the layer has at least this many non-space characters
TEXTLAYER_IMAGE_COVER = 0.7  # image area / page area for a page to count as a scan
_pdf_method = (
    ""  # "pdf+textlayer" when pages were repaired; convert() labels the method
)
_sentence_end = re.compile(r"[.!?:;]['\")\]]?$")


def _nonws_len(text: str) -> int:
    return len(re.sub(r"\s+", "", text))


def _image_only_doc(doc) -> tuple[bool, int, int]:
    """(image_only, hits, sampled): most sampled pages are a full-page image with
    fewer than TEXTLAYER_PAGE_MIN characters of text."""
    import pymupdf

    n = doc.page_count
    idx = (
        sorted({round(i * (n - 1) / 7) for i in range(8)}) if n > 8 else list(range(n))
    )
    hits = 0
    for i in idx:
        page = doc[i]
        area = abs(page.rect) or 1.0
        try:
            imgs = page.get_image_info()
        except Exception:
            imgs = []
        cover = (
            sum(abs(pymupdf.Rect(im["bbox"])) for im in imgs) / area if imgs else 0.0
        )
        if (
            cover >= TEXTLAYER_IMAGE_COVER
            and _nonws_len(page.get_text("text")) < TEXTLAYER_PAGE_MIN
        ):
            hits += 1
    return (bool(idx) and hits / len(idx) >= 0.5), hits, len(idx)


def _textlayer_page_md(page) -> str:
    """One page's text layer as paragraphs. OCR layers often make every LINE its own
    block, so a block is merged into the previous one unless that one ends a sentence
    and this one starts a new one."""
    paras: list[str] = []
    for b in page.get_text("blocks"):
        if len(b) > 6 and b[6] != 0:  # image blocks
            continue
        t = re.sub(
            r"-\n(?=[a-z])", "", b[4]
        )  # re-join words hyphenated at a line break
        t = re.sub(r"\s*\n\s*", " ", t).strip()
        if not t:
            continue
        if paras and not (
            _sentence_end.search(paras[-1]) and (t[0].isupper() or t[0].isdigit())
        ):
            if paras[-1].endswith("-") and t[0].islower():
                paras[-1] = paras[-1][:-1] + t
            else:
                paras[-1] += " " + t
        else:
            paras.append(t)
    return "\n\n".join(paras)


# ---------- per-format extractors ----------
def pdf_to_md(path: str) -> str:
    import pymupdf
    import pymupdf4llm

    global _pdf_method
    _pdf_method = ""
    name = os.path.basename(path)
    try:
        doc = pymupdf.open(path)
        empty, hits, sampled = _image_only_doc(doc)
        if empty:
            doc.close()
            print(
                f"[build_corpus] IMAGE-ONLY {name}: {hits}/{sampled} sampled pages are a "
                "full-page image with no text layer -> OCR route",
                flush=True,
            )
            return ""
        chunks = pymupdf4llm.to_markdown(
            path, page_chunks=True, show_progress=False, use_ocr=False
        )
        out: list[str] = []
        repaired = regained = total = kept = 0
        for i, ch in enumerate(chunks):
            text = ch.get("text", "") if isinstance(ch, dict) else str(ch)
            pno = (
                ((ch.get("metadata") or {}).get("page_number") or (i + 1)) - 1
                if isinstance(ch, dict)
                else i
            )
            if not (0 <= pno < doc.page_count):
                out.append(text)
                continue
            layer, got = _nonws_len(doc[pno].get_text("text")), _nonws_len(text)
            total += layer
            kept += got
            if layer >= TEXTLAYER_PAGE_MIN and got < TEXTLAYER_PAGE_KEEP * layer:
                text = _textlayer_page_md(doc[pno])
                repaired += 1
                regained += layer
            out.append(text)
        doc.close()
        md = "".join(out)
        if repaired:
            _pdf_method = "pdf+textlayer"
            print(
                f"[build_corpus] TEXTLAYER {name}: pymupdf4llm kept {kept} of {total} text-layer "
                f"chars ({kept / max(total, 1):.0%}); {repaired}/{len(chunks)} page(s) replaced by "
                f"their text layer (+{regained} chars)",
                flush=True,
            )
        return clean_md(md)
    except Exception as e:
        # pymupdf4llm raises on some malformed PDFs instead of degrading. Seen
        # on a large scanned PDF: a table is detected but its grid is not, and
        # document_layout.get_table_details then does `grid.h_lines` on None ->
        # AttributeError. The whole document was lost as
        # method="error:AttributeError" with no output.
        #
        # Returning "" instead of raising hands control to convert()'s existing
        # looks_empty() branch, which reuses the OCR cache or runs ocrmypdf --
        # exactly the path a PDF with an unreadable text layer should take. The
        # warning keeps it from being a silent downgrade.
        print(
            f"[build_corpus] WARN: pymupdf4llm failed on {os.path.basename(path)} "
            f"({type(e).__name__}: {e}); falling back to OCR",
            flush=True,
        )
        return ""


def ocr_pdf_to_md(path: str) -> str:
    """OCR a scanned PDF with ocrmypdf (if available) then extract markdown."""
    if not shutil.which("ocrmypdf"):
        return ""
    with tempfile.TemporaryDirectory() as td:
        out = os.path.join(td, "ocr.pdf")
        try:
            subprocess.run(
                ["ocrmypdf", "--force-ocr", "--quiet", path, out],
                check=True,
                timeout=600,
            )
            return pdf_to_md(out)
        except Exception:
            return ""


def docx_to_md(path: str) -> str:
    if shutil.which("pandoc"):
        try:
            return subprocess.run(
                ["pandoc", path, "-t", "gfm"],
                capture_output=True,
                text=True,
                timeout=180,
            ).stdout
        except Exception:
            pass
    try:
        import docx

        return "\n\n".join(
            p.text for p in docx.Document(path).paragraphs if p.text.strip()
        )
    except Exception:
        return ""


def pptx_to_md(path: str) -> str:
    try:
        from pptx import Presentation

        out = []
        for i, s in enumerate(Presentation(path).slides, 1):
            lines = [
                sh.text
                for sh in s.shapes
                if getattr(sh, "has_text_frame", False) and sh.text.strip()
            ]
            if lines:
                out.append(f"## Slide {i}\n\n" + "\n\n".join(lines))
        return "\n\n".join(out)
    except Exception:
        return ""


def epub_to_md(path: str) -> str:
    if shutil.which("pandoc"):
        try:
            return subprocess.run(
                ["pandoc", path, "-t", "gfm"],
                capture_output=True,
                text=True,
                timeout=300,
            ).stdout
        except Exception:
            pass
    try:
        from ebooklib import epub
        from bs4 import BeautifulSoup

        book = epub.read_epub(path)
        parts = []
        for it in book.get_items():
            if it.get_type() == 9:  # DOCUMENT
                parts.append(
                    BeautifulSoup(it.get_content(), "html.parser").get_text("\n")
                )
        return "\n\n".join(parts)
    except Exception:
        return ""


def djvu_to_md(path: str) -> str:
    if shutil.which("djvutxt"):
        try:
            return subprocess.run(
                ["djvutxt", path], capture_output=True, text=True, timeout=300
            ).stdout
        except Exception:
            pass
    return ""


def html_to_md(path: str) -> str:
    if shutil.which("pandoc"):
        try:
            return subprocess.run(
                ["pandoc", path, "-t", "gfm"],
                capture_output=True,
                text=True,
                timeout=120,
            ).stdout
        except Exception:
            pass
    try:
        from bs4 import BeautifulSoup

        return BeautifulSoup(
            open(path, encoding="utf-8", errors="replace").read(), "html.parser"
        ).get_text("\n")
    except Exception:
        return ""


def txt_to_md(path: str) -> str:
    return open(path, encoding="utf-8", errors="replace").read()


# ---------- metadata header ----------
def header(meta: dict, key: str) -> str:
    m = meta.get(key, {})
    title = (m.get("title") or "").strip() or "(untitled)"
    authors = m.get("authors") or []
    if len(authors) > 12:
        authors = authors[:12] + ["et al."]
    lines = [f"# {title}", ""]
    if authors:
        lines.append(f"**Authors:** {'; '.join(authors)}")
    bits = []
    if m.get("year"):
        bits.append(f"**Year:** {m['year']}")
    if m.get("publication"):
        bits.append(f"**Published in:** {m['publication']}")
    if m.get("doi"):
        bits.append(f"**DOI:** {m['doi']}")
    if bits:
        lines.append("  ".join(bits))
    if m.get("tags"):
        lines.append(f"**Tags:** {', '.join(m['tags'][:20])}")
    lines.append(f"**Zotero:** zotero://select/library/items/{key}")
    if m.get("abstract"):
        lines += ["", "> **Abstract.** " + " ".join(m["abstract"].split())]
    lines += ["", "---", ""]
    return "\n".join(lines)


# ---------- worker ----------
CFG = {}  # populated per-process via initializer


def _init(cfg):
    CFG.update(cfg)


def pick_primary(folder: str, want_name: str | None):
    """Choose the best attachment file inside a Zotero storage folder."""
    cands = []
    for fn in os.listdir(folder):
        ext = os.path.splitext(fn)[1].lower()
        if ext in EXT2KIND:
            cands.append((fn, EXT2KIND[ext], os.path.getsize(os.path.join(folder, fn))))
    if not cands:
        return None, None
    if want_name:
        for fn, kind, _ in cands:
            if fn == want_name:
                return os.path.join(folder, fn), kind
    cands.sort(
        key=lambda c: (PRIORITY.index(c[1]), -c[2])
    )  # by kind priority, then size
    fn, kind, _ = cands[0]
    return os.path.join(folder, fn), kind


def convert(key: str):
    meta = CFG["meta"]
    zdir = CFG["zotero"]
    outdir = CFG["out"]
    folder = os.path.join(zdir, "storage", key)
    if not os.path.isdir(folder):
        return (key, "no_folder", "", 0)
    want = (meta.get(key, {}) or {}).get("filename") or None
    src, kind = pick_primary(folder, want)
    if not src:
        return (key, "no_source", "", 0)

    stem = os.path.splitext(os.path.basename(src))[0]
    out_base = f"{key}__{slug(stem)}.md"
    global _ref_doc
    _ref_doc = out_base  # names valve decisions in the build log (one worker = one document at a time)
    # Denylisted output: a superseded duplicate attachment or a webpage snapshot.
    # quarantine_junk.py moves these back out after every build, so converting
    # them is pure waste — on a no-op run it was most of the wall time. Semantics must
    # match quarantine_junk exactly: a bare storage key blocks every file from
    # that attachment, an exact .md filename blocks only that one (so
    # EXAMPLE1__abstract.md stays blocked while a hand-recovered sibling
    # EXAMPLE1__Author_..._recovered.md is never touched by this path at all).
    if key in CFG["deny"] or out_base in CFG["deny"]:
        return (key, "denylist", "", 0)
    out = os.path.join(outdir, out_base)
    # LightRAG moves ingested inputs into __parsed__/ — treat those as done too.
    parsed = os.path.join(outdir, "__parsed__", out_base)
    # ...and an ingest scheduler may park built-but-not-yet-ingested files in a
    # hold dir (--held-dir), which is neither of the above. Without this a held
    # document is rebuilt at the corpus root on every run (at one point, most held
    # files were), and the scheduler then has to discard
    # the duplicate again, which can wedge its queue.
    held = os.path.join(CFG["held"], out_base) if CFG.get("held") else None
    done = (
        out
        if os.path.exists(out)
        else parsed
        if os.path.exists(parsed)
        else held
        if (held and os.path.exists(held))
        else None
    )
    if done == parsed and os.path.getmtime(parsed) < os.path.getmtime(src):
        # When <name> is taken in __parsed__/, LightRAG files the next copy there as
        # <stem>_001.md, _002.md, ... (<stem>_<unix time>.md after 999). It does so
        # after indexing a rebuild whose old document was deleted, and also when a
        # scan skips a rebuild because a processed document still has that name.
        # Either way the build went through a scan, so count the newest copy, or the
        # same rebuild comes back to the corpus root on every run.
        out_stem = out_base[: -len(".md")]
        copies = [
            p
            for p in glob.glob(
                os.path.join(
                    glob.escape(os.path.dirname(parsed)),
                    glob.escape(out_stem) + "_*.md",
                )
            )
            if re.fullmatch(
                r"_(?:\d{3}|\d{10,})", os.path.basename(p)[len(out_stem) : -len(".md")]
            )
        ]
        done = max([parsed] + copies, key=os.path.getmtime)
    if (not CFG["force"]) and done and os.path.getmtime(done) >= os.path.getmtime(src):
        return (key, "skip", done, os.path.getsize(done))

    method = kind
    try:
        if kind == PDF:
            body = pdf_to_md(src)
            if _pdf_method:
                method = _pdf_method
            garbled = looks_garbled(body)
            if garbled:
                print(
                    f"[build_corpus] GARBLED text layer {os.path.basename(src)}: "
                    f"{pua_share(body):.0%} of the text is Private Use Area glyph codes "
                    "-> OCR route",
                    flush=True,
                )
            if looks_empty(body, CFG["min_chars"]) or garbled:
                method = "pdf+ocr"
                # 1) reuse pre-OCR'd text
                pre = (
                    os.path.join(CFG["ocr_cache"], stem + ".pdf.txt")
                    if CFG.get("ocr_cache")
                    else None
                )
                if pre and os.path.exists(pre):
                    body, method = txt_to_md(pre), "pdf+ocr_cache"
                elif CFG["ocr"] in ("auto", "force"):
                    ocr = ocr_pdf_to_md(src)
                    body = ocr or body
                    if not ocr:
                        method = "pdf_needs_ocr"
                # 2) last resort: previously extracted text (e.g. old pdftotext output)
                if looks_empty(body, CFG["min_chars"]) and CFG.get("fallback_md"):
                    old = os.path.join(CFG["fallback_md"], stem + ".md")
                    if os.path.exists(old):
                        body, method = txt_to_md(old), "pdf_legacy_md"
            if garbled and looks_garbled(body):
                # OCR unavailable or failed, and every fallback is the same glyph codes:
                # write the header only, never the garbage.
                body, method = "", "pdf_needs_ocr"
        elif kind == DOCX:
            body = docx_to_md(src)
        elif kind == PPTX:
            body = pptx_to_md(src)
        elif kind == EPUB:
            body = epub_to_md(src)
        elif kind == DJVU:
            body = djvu_to_md(src)
        elif kind == HTML:
            # Prefer Zotero's own text extraction over converting the raw
            # snapshot: .zotero-ft-cache (sibling file in storage/<KEY>/) is far
            # cleaner than any scraper on saved web pages (verified while
            # reviewing quarantined snapshots). Fall back to pandoc/bs4 when the
            # cache is missing or empty.
            ftc = os.path.join(folder, ".zotero-ft-cache")
            body = ""
            if os.path.exists(ftc):
                body, method = txt_to_md(ftc), "html+ftcache"
            if looks_empty(body, CFG["min_chars"]):
                body, method = html_to_md(src), kind
        else:
            body = txt_to_md(src)
    except Exception as e:
        return (key, f"error:{type(e).__name__}", "", 0)

    body = clean_md(
        body or ""
    )  # every path, not just PDF — html snapshots carry base64 image blobs too
    if looks_empty(body, CFG["min_chars"]) and method not in ("pdf_needs_ocr",):
        method += "_empty"
    doc = header(meta, key) + body + "\n"
    tmp = out + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        f.write(doc)
    os.replace(tmp, out)
    if os.path.exists(parsed):
        # LightRAG identifies a document by its file name, so it does not index this
        # rebuild while the document it made from the earlier copy still exists.
        print(
            f"[build_corpus] REBUILT {out_base}: an earlier copy is in __parsed__/; a scan "
            "indexes this one only if LightRAG's document with this name was deleted first",
            flush=True,
        )
    return (key, method, out, len(doc))


# ---------- main ----------
def main():
    ap = argparse.ArgumentParser(
        description="Build a metadata-enriched Markdown corpus from Zotero storage/."
    )
    ap.add_argument(
        "--zotero",
        default=os.path.expanduser("~/Zotero"),
        help="Zotero data dir (contains storage/).",
    )
    ap.add_argument(
        "--out", default=None, help="Output corpus dir (default: <zotero>/rag_corpus)."
    )
    ap.add_argument(
        "--meta",
        default=None,
        help="zotero_metadata.json (default: $ZOTERO_METADATA, else <zotero>/zotero_metadata.json).",
    )
    ap.add_argument("--jobs", type=int, default=max(1, (os.cpu_count() or 4) - 2))
    ap.add_argument(
        "--engine",
        choices=["pymupdf4llm"],
        default="pymupdf4llm",
        help="PDF engine (docling can be added later; pymupdf4llm is fast & CPU-only).",
    )
    ap.add_argument(
        "--ocr",
        choices=["auto", "off", "force"],
        default="auto",
        help="auto/force = run ocrmypdf on image-only PDFs; off = only reuse --ocr-cache-dir.",
    )
    ap.add_argument(
        "--min-chars",
        type=int,
        default=200,
        help="Below this, treat a PDF as image-only.",
    )
    ap.add_argument(
        "--limit", type=int, default=0, help="Process at most N folders (for testing)."
    )
    ap.add_argument(
        "--force", action="store_true", help="Rebuild even if output is up to date."
    )
    ap.add_argument(
        "--denylist",
        default="",
        help="Comma-separated denylist files (bare storage keys or exact .md "
        "filenames) whose output must not be built. Pass the same list "
        "to quarantine_junk.py. Default: none.",
    )
    ap.add_argument(
        "--held-dir",
        default=None,
        help="Optional dir where an ingest scheduler parks built but not yet "
        "ingested files; an up-to-date file there counts as built and is "
        "not rebuilt at the corpus root (default: none).",
    )
    ap.add_argument(
        "--ocr-cache-dir",
        default=None,
        help="Optional dir of pre-OCR'd text, <pdf stem>.pdf.txt, reused for "
        "image-only PDFs before running ocrmypdf (default: none).",
    )
    ap.add_argument(
        "--fallback-md-dir",
        default=None,
        help="Optional dir of previously extracted text, <pdf stem>.md, used "
        "as a last resort for image-only PDFs (default: none).",
    )
    a = ap.parse_args()

    a.out = a.out or os.path.join(a.zotero, "rag_corpus")
    a.meta = (
        a.meta
        or os.environ.get("ZOTERO_METADATA")
        or os.path.join(a.zotero, "zotero_metadata.json")
    )
    os.makedirs(a.out, exist_ok=True)
    meta = json.load(open(a.meta, encoding="utf-8")) if os.path.exists(a.meta) else {}

    storage = os.path.join(a.zotero, "storage")
    keys = sorted(
        d for d in os.listdir(storage) if os.path.isdir(os.path.join(storage, d))
    )
    if a.limit:
        keys = keys[: a.limit]

    deny = load_denylist(a.denylist)

    cfg = {
        "meta": meta,
        "zotero": a.zotero,
        "out": a.out,
        "force": a.force,
        "ocr": a.ocr,
        "min_chars": a.min_chars,
        "deny": deny,
        "held": a.held_dir,
        "ocr_cache": a.ocr_cache_dir,
        "fallback_md": a.fallback_md_dir,
    }
    print(
        f"[build_corpus] {len(keys)} storage folders -> {a.out}  "
        f"(jobs={a.jobs}, ocr={a.ocr}, denylist={len(deny)})",
        flush=True,
    )

    manifest = open(os.path.join(a.out, ".manifest.jsonl"), "w", encoding="utf-8")
    counts, done, t0 = {}, 0, time.time()
    with ProcessPoolExecutor(
        max_workers=a.jobs, initializer=_init, initargs=(cfg,)
    ) as ex:
        futs = {ex.submit(convert, k): k for k in keys}
        for fut in as_completed(futs):
            key, method, out, size = fut.result()
            base = method.split(":")[0]
            counts[base] = counts.get(base, 0) + 1
            manifest.write(
                json.dumps(
                    {
                        "key": key,
                        "method": method,
                        "out": os.path.basename(out),
                        "chars": size,
                    }
                )
                + "\n"
            )
            done += 1
            if done % 50 == 0 or done == len(keys):
                dt = time.time() - t0
                print(
                    f"  {done}/{len(keys)}  ({dt:.0f}s)  "
                    + "  ".join(f"{k}={v}" for k, v in sorted(counts.items())),
                    flush=True,
                )
    manifest.close()
    print("\n[done] methods:", json.dumps(counts, indent=2))
    need = counts.get("pdf_needs_ocr", 0)
    if need:
        print(f"\n⚠  {need} PDFs are image-only and had no OCR available.")
        print(
            "   Install ocrmypdf (`pip install ocrmypdf` + system tesseract) and re-run,"
        )
        print("   or they will be indexed by their metadata header only.")
    print(f"\nCorpus ready: {a.out}")


if __name__ == "__main__":
    main()
